import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, Reference
from openpyxl.cell.cell import MergedCell


def generate_excel_report(
    results,
    output_path="reports/sales_report.xlsx"
):
    """
    Generate a professional Excel sales report.
    """

    # ==========================================
    # CREATE OUTPUT DIRECTORY
    # ==========================================

    output_dir = os.path.dirname(output_path)

    if output_dir:
        os.makedirs(output_dir, exist_ok=True)

    # ==========================================
    # CREATE WORKBOOK
    # ==========================================

    workbook = Workbook()

    # ==========================================
    # SUMMARY SHEET
    # ==========================================

    summary = workbook.active
    summary.title = "Summary"

    summary["A1"] = "AUTOMATED SALES REPORT"

    summary["A1"].font = Font(
        size=18,
        bold=True
    )

    summary.merge_cells("A1:B1")

    summary["A1"].alignment = Alignment(
        horizontal="center"
    )

    summary["A3"] = "Metric"
    summary["B3"] = "Value"

    for cell in ["A3", "B3"]:

        summary[cell].font = Font(
            bold=True,
            color="FFFFFF"
        )

        summary[cell].fill = PatternFill(
            "solid",
            fgColor="1F4E78"
        )

    # ==========================================
    # SUMMARY METRICS
    # ==========================================

    metrics = [
        ("Total Revenue", results["total_revenue"]),
        ("Total Cost", results["total_cost"]),
        ("Total Profit", results["total_profit"]),
        ("Profit Margin", results["profit_margin"] / 100),
        ("Units Sold", results["total_units"]),
        ("Transactions", results["transactions"]),
    ]

    row = 4

    for metric, value in metrics:

        summary.cell(
            row=row,
            column=1,
            value=metric
        )

        summary.cell(
            row=row,
            column=2,
            value=value
        )

        summary.cell(
            row=row,
            column=1
        ).font = Font(bold=True)

        row += 1

    # Currency formatting
    for cell in ["B4", "B5", "B6"]:

        summary[cell].number_format = '"KSh" #,##0.00'

    # Percentage formatting
    summary["B7"].number_format = "0.00%"

    # ==========================================
    # SALES BY PRODUCT
    # ==========================================

    product_sheet = workbook.create_sheet(
        "Sales by Product"
    )

    product_sheet["A1"] = "PRODUCT"
    product_sheet["B1"] = "REVENUE"

    for cell in ["A1", "B1"]:

        product_sheet[cell].font = Font(
            bold=True,
            color="FFFFFF"
        )

        product_sheet[cell].fill = PatternFill(
            "solid",
            fgColor="1F4E78"
        )

    row = 2

    for product, revenue in results[
        "sales_by_product"
    ].items():

        product_sheet.cell(
            row=row,
            column=1,
            value=product
        )

        product_sheet.cell(
            row=row,
            column=2,
            value=float(revenue)
        )

        product_sheet.cell(
            row=row,
            column=2
        ).number_format = '"KSh" #,##0.00'

        row += 1

    # ==========================================
    # PRODUCT CHART
    # ==========================================

    chart = BarChart()

    chart.title = "Revenue by Product"
    chart.y_axis.title = "Revenue (KSh)"
    chart.x_axis.title = "Product"

    data_reference = Reference(
        product_sheet,
        min_col=2,
        min_row=1,
        max_row=row - 1
    )

    categories = Reference(
        product_sheet,
        min_col=1,
        min_row=2,
        max_row=row - 1
    )

    chart.add_data(
        data_reference,
        titles_from_data=True
    )

    chart.set_categories(categories)

    chart.height = 8
    chart.width = 15

    product_sheet.add_chart(
        chart,
        "D2"
    )

    # ==========================================
    # SALES BY CATEGORY
    # ==========================================

    category_sheet = workbook.create_sheet(
        "Sales by Category"
    )

    category_sheet["A1"] = "CATEGORY"
    category_sheet["B1"] = "REVENUE"

    for cell in ["A1", "B1"]:

        category_sheet[cell].font = Font(
            bold=True,
            color="FFFFFF"
        )

        category_sheet[cell].fill = PatternFill(
            "solid",
            fgColor="1F4E78"
        )

    row = 2

    for category, revenue in results[
        "sales_by_category"
    ].items():

        category_sheet.cell(
            row=row,
            column=1,
            value=category
        )

        category_sheet.cell(
            row=row,
            column=2,
            value=float(revenue)
        )

        category_sheet.cell(
            row=row,
            column=2
        ).number_format = '"KSh" #,##0.00'

        row += 1

    # ==========================================
    # RAW DATA
    # ==========================================

    raw_sheet = workbook.create_sheet(
        "Raw Data"
    )

    data = results["data"]

    # Headers
    for column_index, column_name in enumerate(
        data.columns,
        start=1
    ):

        cell = raw_sheet.cell(
            row=1,
            column=column_index,
            value=column_name
        )

        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )

        cell.fill = PatternFill(
            "solid",
            fgColor="1F4E78"
        )

    # Data
    for row_index, row_data in enumerate(
        data.itertuples(index=False),
        start=2
    ):

        for column_index, value in enumerate(
            row_data,
            start=1
        ):

            raw_sheet.cell(
                row=row_index,
                column=column_index,
                value=value
            )

    # ==========================================
    # AUTO-SIZE COLUMNS
    # ==========================================

    for sheet in workbook.worksheets:

        for column_cells in sheet.iter_cols():

            max_length = 0

            for cell in column_cells:

                # Skip merged cells
                if isinstance(cell, MergedCell):
                    continue

                if cell.value is not None:

                    length = len(
                        str(cell.value)
                    )

                    if length > max_length:
                        max_length = length

            # Get the column letter safely
            if column_cells:

                first_cell = column_cells[0]

                if isinstance(
                    first_cell,
                    MergedCell
                ):
                    continue

                column_letter = (
                    first_cell.column_letter
                )

                sheet.column_dimensions[
                    column_letter
                ].width = min(
                    max_length + 2,
                    40
                )

    # ==========================================
    # SAVE WORKBOOK
    # ==========================================

    workbook.save(output_path)

    return output_path